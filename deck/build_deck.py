"""Build the AgentCore client deck as an image-only PPTX.

Each slide is authored as a 1920x1080 HTML page, rendered to PNG with headless
Chrome (2x scale) and placed full-bleed on a 16:9 PowerPoint slide. The slide
text is also written into the speaker notes so presenters can search it.
"""
from __future__ import annotations

import json
import math
import pathlib
import re
import subprocess
import sys

from pptx import Presentation
from pptx.util import Inches

ROOT = pathlib.Path(__file__).resolve().parent
OUT = ROOT / "out"
HTML_DIR = OUT / "html"
PNG_DIR = OUT / "png"
PPTX_PATH = ROOT / "AgentCore_Client_Deck.pptx"
CHROME = pathlib.Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")
ICON_DIR = pathlib.Path(
    r"C:\Users\81691\Downloads\Agentcore_final\agentcore\src\frontend\node_modules\lucide-react\dist\esm\icons"
)

# ----------------------------------------------------------------------------
# Palette (dataviz reference palette, dark steps, validated on #0b1220)
# ----------------------------------------------------------------------------
C = {
    "surface": "#0b1220",
    "text": "#f0efec",
    "text2": "#c3c2b7",
    "muted": "#8b8f98",
    "blue": "#3987e5",
    "blue_l": "#6da7ec",
    "orange": "#d95926",
    "aqua": "#199e70",
    "yellow": "#c98500",
    "magenta": "#d55181",
    "violet": "#9085e9",
    "good": "#0ca30c",
    "warning": "#fab219",
    "serious": "#ec835a",
    "critical": "#d03b3b",
}

# ----------------------------------------------------------------------------
# Icons (lucide, extracted from the frontend's node_modules)
# ----------------------------------------------------------------------------
_icon_cache: dict[str, list] = {}
_missing_icons: set[str] = set()


def _icon_nodes(name: str) -> list:
    if name in _icon_cache:
        return _icon_cache[name]
    path = ICON_DIR / f"{name}.js"
    if not path.exists():
        _missing_icons.add(name)
        path = ICON_DIR / "circle.js"
    src = path.read_text(encoding="utf-8")
    alias = re.search(r"export \{ default \} from '\./([\w-]+)\.js'", src)
    if alias:  # deprecated alias files re-export the renamed icon
        nodes = _icon_nodes(alias.group(1))
        _icon_cache[name] = nodes
        return nodes
    m = re.search(r"const __iconNode = (\[.*?\]);\s*\nconst", src, re.S)
    arr = re.sub(r'([{,\[]\s*)([A-Za-z_][\w-]*)\s*:', r'\1"\2":', m.group(1))
    nodes = json.loads(arr)
    _icon_cache[name] = nodes
    return nodes


def _kebab(k: str) -> str:
    return re.sub(r"([A-Z])", lambda m: "-" + m.group(1).lower(), k)


def icon(name: str, size: int = 28, color: str = "currentColor", sw: float = 1.75, cls: str = "") -> str:
    parts = []
    for tag, attrs in _icon_nodes(name):
        a = " ".join(f'{_kebab(k)}="{v}"' for k, v in attrs.items() if k != "key")
        parts.append(f"<{tag} {a}/>")
    return (
        f'<svg class="{cls}" width="{size}" height="{size}" viewBox="0 0 24 24" fill="none" '
        f'stroke="{color}" stroke-width="{sw}" stroke-linecap="round" stroke-linejoin="round">'
        f'{"".join(parts)}</svg>'
    )


def logo(size: int = 44) -> str:
    """AgentCore mark: rounded square with a small node graph."""
    return f"""<svg width="{size}" height="{size}" viewBox="0 0 48 48" fill="none">
  <defs><linearGradient id="lg" x1="0" y1="0" x2="1" y2="1">
    <stop offset="0" stop-color="{C['blue']}"/><stop offset="1" stop-color="{C['violet']}"/></linearGradient></defs>
  <rect x="2" y="2" width="44" height="44" rx="12" fill="url(#lg)"/>
  <path d="M15 16h8M23 16l10 8M33 24l-10 8M15 32h8" stroke="#fff" stroke-width="2.4" stroke-linecap="round"/>
  <circle cx="14" cy="16" r="4" fill="#fff"/><circle cx="14" cy="32" r="4" fill="#fff"/>
  <circle cx="34" cy="24" r="4.5" fill="#0b1220" stroke="#fff" stroke-width="2.4"/>
</svg>"""


# ----------------------------------------------------------------------------
# Shared CSS
# ----------------------------------------------------------------------------
CSS = f"""
*{{box-sizing:border-box;margin:0;padding:0}}
html,body{{width:1920px;height:1080px;overflow:hidden}}
body{{font-family:"Segoe UI Variable Display","Segoe UI",system-ui,sans-serif;background:{C['surface']};color:{C['text']};position:relative;-webkit-font-smoothing:antialiased}}
.bg{{position:absolute;inset:0;background:
  radial-gradient(1300px 760px at 88% -12%, rgba(57,135,229,.22), transparent 62%),
  radial-gradient(1000px 640px at -8% 112%, rgba(144,133,233,.16), transparent 62%),
  {C['surface']}}}
.grid{{position:absolute;inset:0;background-image:linear-gradient(rgba(255,255,255,.035) 1px,transparent 1px),linear-gradient(90deg,rgba(255,255,255,.035) 1px,transparent 1px);background-size:80px 80px;-webkit-mask-image:radial-gradient(ellipse at 50% 45%, #000 35%, transparent 85%)}}
.slide{{position:absolute;inset:0}}
.head{{position:absolute;left:96px;right:96px;top:64px}}
.eyebrow{{font-size:20px;letter-spacing:.2em;text-transform:uppercase;color:{C['blue_l']};font-weight:600;margin-bottom:14px}}
h1.title{{font-size:56px;font-weight:600;line-height:1.1;letter-spacing:-.012em}}
.sub{{font-size:26px;color:{C['text2']};margin-top:14px;max-width:1500px;line-height:1.35}}
.body{{position:absolute;left:96px;right:96px;top:288px;bottom:104px}}
.footer{{position:absolute;left:96px;right:96px;bottom:38px;display:flex;justify-content:space-between;align-items:center;color:{C['muted']};font-size:18px}}
.brand{{display:flex;align-items:center;gap:12px;font-weight:600;color:{C['text2']};font-size:20px;letter-spacing:.01em}}
.card{{background:rgba(255,255,255,.045);border:1px solid rgba(255,255,255,.10);border-radius:22px;padding:28px}}
.ic{{display:inline-flex;align-items:center;justify-content:center;border-radius:14px;flex:none}}
ul.dots{{list-style:none}}
ul.dots li{{position:relative;padding-left:26px;font-size:22px;color:{C['text2']};line-height:1.35;margin-top:10px}}
ul.dots li::before{{content:"";position:absolute;left:4px;top:13px;width:8px;height:8px;border-radius:50%;background:var(--c,{C['blue']})}}
.chip{{display:inline-flex;align-items:center;gap:10px;padding:10px 18px;border-radius:999px;border:1px solid rgba(255,255,255,.14);background:rgba(255,255,255,.05);font-size:21px;color:{C['text']};white-space:nowrap}}
.chip svg{{flex:none}}
.pill{{display:inline-flex;align-items:center;gap:8px;padding:6px 14px;border-radius:999px;font-size:18px;font-weight:600;letter-spacing:.02em}}
.muted{{color:{C['muted']}}}
.t2{{color:{C['text2']}}}
"""


def frame(n: int, eyebrow: str, title: str, sub: str, body: str, extra_css: str = "") -> str:
    sub_html = f'<p class="sub">{sub}</p>' if sub else ""
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{CSS}{extra_css}</style></head><body>
<div class="bg"></div><div class="grid"></div>
<div class="slide">
  <div class="head"><div class="eyebrow">{eyebrow}</div><h1 class="title">{title}</h1>{sub_html}</div>
  <div class="body">{body}</div>
  <div class="footer"><div class="brand">{logo(30)}<span>AgentCore</span></div><div>Enterprise AI agent platform&nbsp;&nbsp;·&nbsp;&nbsp;{n:02d}</div></div>
</div></body></html>"""


def bare(body: str, extra_css: str = "") -> str:
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{CSS}{extra_css}</style></head><body>
<div class="bg"></div><div class="grid"></div>{body}</body></html>"""


def ic(name: str, color: str, size: int = 56, isize: int = 28, radius: int = 14) -> str:
    return (
        f'<span class="ic" style="width:{size}px;height:{size}px;border-radius:{radius}px;'
        f'background:{color}22;border:1px solid {color}66;color:{color}">{icon(name, isize)}</span>'
    )


def network_svg(w: int, h: int, opacity: float = 0.5) -> str:
    pts = [(120, 160), (340, 90), (520, 240), (760, 120), (930, 300), (1120, 160), (250, 420),
           (480, 520), (700, 420), (960, 560), (1180, 420), (380, 700), (640, 760), (880, 700), (1100, 820)]
    edges = [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (0, 6), (6, 7), (2, 7), (7, 8), (8, 4), (8, 9), (9, 10), (5, 10),
             (6, 11), (11, 12), (7, 12), (12, 13), (13, 9), (13, 14), (10, 14)]
    lines = "".join(
        f'<line x1="{pts[a][0]}" y1="{pts[a][1]}" x2="{pts[b][0]}" y2="{pts[b][1]}" stroke="{C["blue_l"]}" stroke-opacity=".35" stroke-width="1.5"/>'
        for a, b in edges
    )
    dots = "".join(
        f'<circle cx="{x}" cy="{y}" r="{6 if i % 3 else 9}" fill="{C["blue"] if i % 4 else C["violet"]}" fill-opacity=".9"/>'
        for i, (x, y) in enumerate(pts)
    )
    return f'<svg width="{w}" height="{h}" viewBox="0 0 1300 900" style="opacity:{opacity}">{lines}{dots}</svg>'


# ----------------------------------------------------------------------------
# Slides
# ----------------------------------------------------------------------------
SLIDES: list[tuple[str, str]] = []  # (html, notes)


def add(html: str, notes: str) -> None:
    SLIDES.append((html, notes))


# 1 — Cover -------------------------------------------------------------------
add(
    bare(f"""
<div style="position:absolute;right:-40px;top:60px">{network_svg(1200, 830, .55)}</div>
<div style="position:absolute;left:120px;top:150px;display:flex;align-items:center;gap:26px">
  {logo(96)}<span style="font-size:72px;font-weight:700;letter-spacing:-.02em">AgentCore</span>
</div>
<div style="position:absolute;left:120px;top:330px;max-width:1150px">
  <div style="font-size:92px;font-weight:700;line-height:1.04;letter-spacing:-.02em">Build, govern and run<br/>AI agents at enterprise scale</div>
  <div style="font-size:32px;color:{C['text2']};margin-top:36px;line-height:1.4;max-width:980px">One platform for the visual design of agents, their connection to enterprise systems, governance and production operations.</div>
</div>
<div style="position:absolute;left:120px;bottom:96px;display:flex;gap:16px">
  <span class="chip">{icon('workflow',22,C['blue_l'])}Build</span>
  <span class="chip">{icon('plug',22,C['blue_l'])}Connect</span>
  <span class="chip">{icon('shield',22,C['blue_l'])}Govern</span>
  <span class="chip">{icon('rocket',22,C['blue_l'])}Run</span>
</div>
<div style="position:absolute;right:120px;bottom:96px;color:{C['muted']};font-size:22px;text-align:right">Client briefing<br/>October 2026</div>
"""),
    "AgentCore: build, govern and run AI agents at enterprise scale. One platform for the visual design of agents, connection to enterprise systems, governance and production operations. Four pillars: Build, Connect, Govern, Run.",
)

# 2 — Problem -----------------------------------------------------------------
steps = ["Idea", "Prototype", "Pilot", "Production"]
step_html = ""
for i, s in enumerate(steps):
    done = i < 3
    col = C["blue"] if done else C["critical"]
    step_html += f"""
    <div style="display:flex;flex-direction:column;align-items:center;gap:14px;width:300px">
      <div style="width:64px;height:64px;border-radius:50%;display:flex;align-items:center;justify-content:center;background:{col}22;border:2px solid {col};color:{col}">
        {icon('check' if done else 'x', 30, col, 2.5)}</div>
      <div style="font-size:26px;font-weight:600;color:{C['text'] if done else C['serious']}">{s}</div>
    </div>"""
    if i < 3:
        dash = "none" if i < 2 else "14 12"
        step_html += f'<div style="flex:1;height:0;border-top:3px {"solid" if i < 2 else "dashed"} {C["blue"] if i < 2 else C["critical"]};margin-top:31px;opacity:.8"></div>'

blockers = [
    ("boxes", C["blue"], "Fragmented tooling", "Every team picks its own models, vector stores and prompts. Nothing is reusable or comparable."),
    ("shield-alert", C["orange"], "No governance", "Nobody can say which model was approved, who may send mail on the company's behalf, or what changed."),
    ("plug-zap", C["aqua"], "Integration gaps", "Mailboxes, SharePoint, databases and legacy systems stay out of reach of the agent."),
    ("eye-off", C["violet"], "No visibility", "Cost, latency and failures are discovered after the fact, if at all."),
]
cards = "".join(
    f"""<div class="card" style="flex:1;display:flex;flex-direction:column;gap:18px;padding:32px">
      {ic(n, c, 64, 32)}
      <div style="font-size:32px;font-weight:600">{t}</div>
      <div style="font-size:24px;color:{C['text2']};line-height:1.45">{d}</div></div>"""
    for n, c, t, d in blockers
)
add(
    frame(2, "Why now", "Agents are easy to prototype and hard to run in production",
          "Most initiatives stall between the pilot and the production stage. The reasons are rarely about the model.",
          f"""<div style="display:flex;flex-direction:column;height:100%"><div style="display:flex;align-items:flex-start;padding:0 60px;margin-top:4px">{step_html}</div>
<div style="display:flex;gap:28px;margin-top:56px;flex:1">{cards}</div></div>"""),
    "Agents are easy to prototype and hard to run in production. Four blockers: fragmented tooling, no governance, integration gaps, no visibility.",
)

# 3 — One platform, four pillars ----------------------------------------------
pillars = [
    ("workflow", C["blue"], "Build", ["Visual agent studio", "Curated component library", "Playground for testing", "Multi-agent patterns", "Templates and starter projects"]),
    ("plug", C["aqua"], "Connect", ["Any LLM provider through one gateway", "Microsoft 365 and databases", "Knowledge bases, vectors and graphs", "MCP tool servers", "Scheduled and event-driven triggers"]),
    ("shield", C["violet"], "Govern", ["Roles and workspaces", "Approval workflows", "Guardrails and human review", "Audit trails", "Cost limits and timeouts"]),
    ("rocket", C["orange"], "Run", ["Releases across DEV, UAT and PROD", "APIs, chat widget and Teams", "Monitoring and tracing", "Evaluation datasets", "Multi-region deployment"]),
]
pill_html = "".join(
    f"""<div class="card" style="flex:1;padding:36px 32px;--c:{c};border-top:4px solid {c}">
      <div style="display:flex;align-items:center;gap:18px">{ic(n, c, 64, 32)}<div style="font-size:38px;font-weight:700">{t}</div></div>
      <ul class="dots" style="margin-top:22px">{"".join(f"<li>{b}</li>" for b in bs)}</ul></div>"""
    for n, c, t, bs in pillars
)
add(
    frame(3, "What AgentCore is", "One platform from first prototype to governed production",
          "AgentCore brings the full agent lifecycle under one login, one catalogue and one audit trail.",
          f"""<div style="display:flex;gap:28px;height:560px">{pill_html}</div>
<div style="display:flex;justify-content:center;gap:18px;margin-top:40px">
  <span class="chip">{icon('log-in',22,C['blue_l'])}One login</span>
  <span class="chip">{icon('library',22,C['blue_l'])}One catalogue of agents, models, tools and guardrails</span>
  <span class="chip">{icon('scroll-text',22,C['blue_l'])}One audit trail</span></div>""",
          extra_css="ul.dots li{font-size:25px;margin-top:14px}"),
    "One platform from first prototype to governed production. Build: visual studio, component library, playground, multi-agent. Connect: LLM gateway, Microsoft 365, databases, knowledge, MCP. Govern: roles, approvals, guardrails, audit. Run: releases, APIs, widget, Teams, schedules, monitoring.",
)

# 4 — Architecture ------------------------------------------------------------
layers = [
    ("Experience", C["blue"], "monitor", [
        ("layout-dashboard", "Web studio"), ("play", "Playground"), ("message-square", "Embeddable chat widget"),
        ("users", "Microsoft Teams"), ("bot", "Orchestrator chat"), ("braces", "REST and OpenAI-compatible API"), ("share-2", "Agent-to-Agent protocol")]),
    ("Platform core", C["violet"], "cpu", [
        ("workflow", "Agent runtime"), ("folder-kanban", "Workspaces and RBAC"), ("badge-check", "Approvals and releases"),
        ("calendar-clock", "Scheduler and triggers"), ("user-check", "Human-in-the-loop"), ("flask-conical", "Evaluation"),
        ("coins", "Cost limits"), ("activity", "Observability")]),
    ("Services", C["aqua"], "server", [
        ("brain", "Model gateway and registry"), ("shield-check", "Guardrails service"), ("wrench", "MCP tool broker"),
        ("database", "Vector store service"), ("network", "GraphRAG service"), ("globe", "Region gateway")]),
    ("Data and cloud", C["yellow"], "cloud", [
        ("database", "PostgreSQL"), ("zap", "Redis"), ("mail", "RabbitMQ"), ("hard-drive", "Blob storage"),
        ("key-round", "Azure Key Vault"), ("fingerprint", "Entra ID"), ("scan-search", "Langfuse"), ("bar-chart-3", "Prometheus and Grafana")]),
]
rows = ""
for name, col, icn, chips in layers:
    chip_html = "".join(f'<span class="chip" style="font-size:20px;padding:9px 16px">{icon(c_i,20,col)}{c_t}</span>' for c_i, c_t in chips)
    rows += f"""
    <div style="display:flex;align-items:center;gap:28px;background:rgba(255,255,255,.035);border:1px solid rgba(255,255,255,.09);border-left:6px solid {col};border-radius:18px;padding:22px 28px">
      <div style="width:260px;display:flex;align-items:center;gap:14px;flex:none">{ic(icn,col,52,26,12)}<div style="font-size:26px;font-weight:600">{name}</div></div>
      <div style="display:flex;flex-wrap:wrap;gap:12px">{chip_html}</div></div>"""
add(
    frame(4, "Architecture", "A modular platform built from proven building blocks",
          "A web studio and APIs on top, a FastAPI core, independent microservices and standard Azure data services underneath.",
          f'<div style="display:flex;flex-direction:column;gap:18px">{rows}</div>'),
    "Architecture layers. Experience: web studio, playground, embeddable chat widget, Teams, orchestrator chat, REST/OpenAI-compatible API, A2A. Platform core: agent runtime, workspaces and RBAC, approvals and releases, scheduler and triggers, HITL, evaluation, cost limits, observability. Services: model gateway, guardrails, MCP tool broker, vector store (Pinecone), GraphRAG (Neo4j), region gateway. Data and cloud: PostgreSQL, Redis, RabbitMQ, Blob storage, Key Vault, Entra ID, Langfuse, Prometheus/Grafana.",
)

# 5 — Build: visual studio -----------------------------------------------------
nodes = [
    (40, 232, "message-square", C["blue"], "Chat input", "Email body and attachments"),
    (300, 100, "scroll-text", C["violet"], "Prompt", "Underwriting instructions"),
    (300, 364, "brain", C["aqua"], "Registry LLM", "Approved model, rate-limited"),
    (560, 232, "database", C["orange"], "Database tool", "Natural language to SQL"),
    (820, 100, "split", C["magenta"], "Conditional router", "Proceed, clarify or refer"),
    (820, 364, "message-square", C["blue"], "Chat output", "Decision and quote draft"),
]
W, Hn = 220, 110
edges = [(0, 1), (0, 2), (1, 3), (2, 3), (3, 4), (3, 5)]


def bez(a, b):
    x1, y1 = nodes[a][0] + W, nodes[a][1] + Hn / 2
    x2, y2 = nodes[b][0], nodes[b][1] + Hn / 2
    dx = (x2 - x1) * 0.5
    return f'<path d="M{x1},{y1} C{x1+dx},{y1} {x2-dx},{y2} {x2},{y2}" stroke="{C["blue_l"]}" stroke-opacity=".7" stroke-width="3" fill="none"/>'


edge_svg = "".join(bez(a, b) for a, b in edges)
node_html = "".join(
    f"""<div style="position:absolute;left:{x}px;top:{y}px;width:{W}px;height:{Hn}px;background:#121a2c;border:1px solid rgba(255,255,255,.14);border-top:4px solid {c};border-radius:14px;padding:14px 16px;box-shadow:0 10px 30px rgba(0,0,0,.35)">
      <div style="display:flex;align-items:center;gap:10px;font-size:20px;font-weight:600">{icon(i,20,c)}{t}</div>
      <div style="font-size:16px;color:{C['muted']};margin-top:8px;line-height:1.3">{d}</div>
      <span style="position:absolute;left:-7px;top:50%;width:12px;height:12px;border-radius:50%;background:{C['blue_l']};transform:translateY(-50%)"></span>
      <span style="position:absolute;right:-7px;top:50%;width:12px;height:12px;border-radius:50%;background:{C['blue_l']};transform:translateY(-50%)"></span></div>"""
    for x, y, i, c, t, d in nodes
)
canvas = f"""<div style="position:relative;width:1080px;height:580px;border-radius:22px;border:1px solid rgba(255,255,255,.10);background:
  linear-gradient(rgba(255,255,255,.05) 1px,transparent 1px) 0 0/40px 40px,linear-gradient(90deg,rgba(255,255,255,.05) 1px,transparent 1px) 0 0/40px 40px,#0e1526;overflow:hidden">
  <svg style="position:absolute;inset:0" width="1080" height="580">{edge_svg}</svg>{node_html}
  <div style="position:absolute;left:20px;bottom:18px;display:flex;gap:10px">
    <span class="pill" style="background:{C['good']}22;color:#7ee07e;border:1px solid {C['good']}66">{icon('play',16,'#7ee07e')}Run in playground</span>
    <span class="pill" style="background:rgba(255,255,255,.06);color:{C['text2']};border:1px solid rgba(255,255,255,.14)">{icon('save',16)}Save version</span></div></div>"""
cats = [
    ("bot", "Agents"), ("brain", "Models"), ("wrench", "Tools"), ("database", "Vector stores"), ("network", "GraphRAG"),
    ("shield-check", "Guardrails"), ("user-check", "Human-in-the-loop"), ("scan-text", "OCR"), ("git-branch", "Logic"),
    ("sliders-horizontal", "Processing"), ("calendar-clock", "Triggers"), ("plug", "MCP"), ("code-2", "Custom components"),
]
cat_html = "".join(f'<span class="chip" style="font-size:20px">{icon(i,20,C["blue_l"])}{t}</span>' for i, t in cats)
tools = ["Database connector", "Natural-language SQL", "Data visualiser", "Outlook mail", "SharePoint documents", "Web search", "API request", "Smart router", "Supervisor and collaborative agents"]
tool_html = "".join(f"<li>{t}</li>" for t in tools)
add(
    frame(5, "Build", "Design agents visually, test them instantly",
          "Business and technical teams assemble agents on a canvas from a curated library of enterprise-ready components.",
          f"""<div style="display:flex;gap:40px;align-items:flex-start">{canvas}
<div style="flex:1">
  <div style="font-size:28px;font-weight:600;margin-bottom:16px">Component library</div>
  <div style="display:flex;flex-wrap:wrap;gap:10px">{cat_html}</div>
  <div style="font-size:24px;font-weight:600;margin-top:28px;color:{C['text2']}">Ready-made tools include</div>
  <ul class="dots" style="--c:{C['aqua']};margin-top:4px;columns:1">{tool_html}</ul></div></div>"""),
    "Build: a visual canvas for designing agents. Component categories: agents, models, tools, vector stores, GraphRAG, guardrails, human-in-the-loop, OCR, logic, processing, triggers, MCP, custom components. Ready tools: database connector, NL-SQL, data visualiser, Outlook, SharePoint, web search, API request, smart router, supervisor and collaborative agents. Playground to run and save versions.",
)

# 6 — Models: gateway and registry --------------------------------------------
providers = ["OpenAI", "Azure OpenAI", "Anthropic", "Google Gemini", "Google Vertex AI", "Groq", "Any OpenAI-compatible endpoint"]
cx, cy, R = 480, 300, 240
spokes = ""
lines = ""
n = len(providers)
for i, p in enumerate(providers):
    ang = -math.pi / 2 + i * 2 * math.pi / n
    px, py = cx + R * math.cos(ang) * 1.45, cy + R * math.sin(ang)
    lines += f'<line x1="{cx}" y1="{cy}" x2="{px:.0f}" y2="{py:.0f}" stroke="{C["blue_l"]}" stroke-opacity=".45" stroke-width="2"/>'
    spokes += f'<div class="chip" style="position:absolute;left:{px:.0f}px;top:{py:.0f}px;transform:translate(-50%,-50%);font-size:21px;background:#121a2c">{icon("sparkles",20,C["aqua"])}{p}</div>'
hub = f"""<div style="position:relative;width:960px;height:600px">
  <svg style="position:absolute;inset:0" width="960" height="600">{lines}</svg>{spokes}
  <div style="position:absolute;left:{cx}px;top:{cy}px;transform:translate(-50%,-50%);width:260px;height:260px;border-radius:50%;background:radial-gradient(circle at 50% 40%, {C['blue']}55, #121a2c 70%);border:2px solid {C['blue']};display:flex;flex-direction:column;align-items:center;justify-content:center;gap:10px;text-align:center;box-shadow:0 0 60px {C['blue']}44">
    {icon('brain',52,C['text'])}<div style="font-size:26px;font-weight:700;line-height:1.15">Model gateway</div><div style="font-size:17px;color:{C['text2']};padding:0 24px">one OpenAI-compatible endpoint for chat and embeddings</div></div></div>"""
feats = [
    ("badge-check", "Approval-enforced promotion", "A model is usable by agents only after it has been approved and promoted."),
    ("gauge", "Per-model rate limits", "Request and token limits are set centrally per model, not per developer."),
    ("key-round", "Keys never leave the vault", "Provider API keys are stored encrypted and are invisible to agent builders."),
    ("scroll-text", "Model audit log", "Every registry change and promotion is recorded with who, what and when."),
    ("plug-zap", "Test connection", "Models are verified against the provider before they reach the catalogue."),
]
feat_html = "".join(
    f"""<div style="display:flex;gap:18px;align-items:flex-start">{ic(i,C['aqua'],52,26,12)}
      <div><div style="font-size:24px;font-weight:600">{t}</div><div style="font-size:19px;color:{C['text2']};line-height:1.35;margin-top:4px">{d}</div></div></div>"""
    for i, t, d in feats
)
add(
    frame(6, "Connect · Models", "Every model behind one governed gateway",
          "One endpoint for every agent. The registry decides which models are allowed, at what rate and for whom.",
          f"""<div style="display:flex;gap:40px;align-items:center">{hub}
<div class="card" style="flex:1;display:flex;flex-direction:column;gap:24px;padding:32px">{feat_html}</div></div>"""),
    "Models: one OpenAI-compatible gateway for chat and embeddings across OpenAI, Azure OpenAI, Anthropic, Google Gemini, Google Vertex AI, Groq and any OpenAI-compatible endpoint. Registry: approval-enforced promotion, per-model rate limits, keys stored encrypted, model audit log, test connection.",
)

# 7 — Connect: enterprise systems and knowledge --------------------------------
systems = [
    ("mail", "Outlook mail", "Read, triage and reply via Microsoft Graph"),
    ("folder-open", "SharePoint", "Documents and libraries"),
    ("users", "Microsoft Teams", "Agents as Teams bots"),
    ("database", "PostgreSQL, MySQL, Oracle, SQL Server", "Governed database connectors"),
    ("hard-drive", "Azure Blob storage", "Files and attachments"),
    ("wrench", "MCP tool servers", "Any tool over the Model Context Protocol"),
    ("globe", "Web search and REST APIs", "External data on demand"),
    ("calendar-clock", "Schedules and folder monitors", "Agents that start themselves"),
]
knowledge = [
    ("book-open", "Knowledge bases", "Curated document collections per workspace"),
    ("database", "Pinecone and Azure AI Search", "Vector retrieval at scale"),
    ("network", "GraphRAG on Neo4j", "Entity extraction, resolution and community detection"),
    ("search", "Semantic search", "Across knowledge bases from any agent"),
    ("scan-text", "OCR and document intelligence", "Scanned PDFs become structured fields"),
    ("history", "Long-term memory", "Facts and summaries that persist across sessions"),
]


def conn_panel(title, icn, col, items):
    rows_ = "".join(
        f"""<div style="display:flex;gap:16px;align-items:flex-start">{ic(i,col,50,25,12)}
          <div><div style="font-size:24px;font-weight:600;line-height:1.2">{t}</div><div style="font-size:19px;color:{C['text2']};margin-top:4px;line-height:1.3">{d}</div></div></div>"""
        for i, t, d in items
    )
    return f"""<div class="card" style="flex:1;padding:34px 36px;border-top:4px solid {col}">
      <div style="display:flex;align-items:center;gap:16px;margin-bottom:30px">{ic(icn,col,56,28)}<div style="font-size:32px;font-weight:700">{title}</div></div>
      <div style="display:grid;grid-template-columns:1fr 1fr;gap:34px 28px">{rows_}</div></div>"""


add(
    frame(7, "Connect · Systems and knowledge", "Agents that reach the systems your people already use",
          "Microsoft 365, databases and tool servers on one side. Documents, vectors and graphs on the other.",
          f"""<div style="display:flex;gap:28px">{conn_panel('Enterprise systems','plug',C['blue'],systems)}{conn_panel('Knowledge',"book-open",C['violet'],knowledge)}</div>"""),
    "Connect: Outlook mail, SharePoint, Teams, PostgreSQL/MySQL/Oracle/SQL Server, Azure Blob, MCP tool servers, web search and REST APIs, schedules and folder monitors. Knowledge: knowledge bases, Pinecone and Azure AI Search, GraphRAG on Neo4j, semantic search, OCR and document intelligence, long-term memory.",
)

# 8 — Govern -------------------------------------------------------------------
gov = [
    ("users", C["blue"], "Role-based access", "Organisations, departments and projects as workspaces. Roles and fine-grained permissions decide who builds, approves and runs.", ["Organisation", "Department", "Project"]),
    ("badge-check", C["violet"], "Approval workflows", "New models, MCP servers and production deployments wait for a named approver.", ["Models", "MCP servers", "PROD releases"]),
    ("shield-check", C["aqua"], "Guardrails catalogue", "Versioned policies on NeMo Guardrails with PII detection. Promote, demote and apply them to any agent.", ["NeMo Guardrails", "Presidio PII", "Versioned"]),
    ("user-check", C["orange"], "Human-in-the-loop", "Approval steps inside the agent flow. The agent pauses until a person decides.", ["Approval component", "Review queue"]),
    ("scroll-text", C["magenta"], "Audit trails", "Model, MCP and guardrail execution logs. Every run is traceable to inputs, documents and decisions.", ["Model log", "MCP log", "Guardrail log"]),
    ("package-check", C["yellow"], "Dependency and limits", "Approved packages, execution timeouts and cost thresholds with alerts.", ["Packages", "Timeouts", "Cost alerts"]),
]
gov_html = "".join(
    f"""<div class="card" style="padding:30px;display:flex;flex-direction:column;gap:16px;border-top:4px solid {c}">
      <div style="display:flex;align-items:center;gap:16px">{ic(i,c,58,30)}<div style="font-size:28px;font-weight:600">{t}</div></div>
      <div style="font-size:23px;color:{C['text2']};line-height:1.45">{d}</div>
      <div style="margin-top:auto;display:flex;flex-wrap:wrap;gap:8px">{"".join(f'<span class="chip" style="font-size:18px;padding:6px 14px;color:{C["text2"]}">{tg}</span>' for tg in tags)}</div></div>"""
    for i, c, t, d, tags in gov
)
add(
    frame(8, "Govern", "Governance is built into the platform, not added later",
          "Controls apply to every agent automatically: who may do what, which models and tools are allowed, and what gets recorded.",
          f'<div style="display:grid;grid-template-columns:repeat(3,1fr);gap:26px;height:100%">{gov_html}</div>'),
    "Govern: role-based access with organisation/department/project workspaces; approval workflows for models, MCP servers and PROD deployments; guardrails catalogue on NeMo Guardrails with Presidio PII; human-in-the-loop; audit trails; dependency governance, timeouts and cost limits.",
)

# 9 — Run ----------------------------------------------------------------------
stages = [("pencil-ruler", "Design", "Canvas and playground", C["blue"]), ("flask-conical", "DEV", "Iterate freely", C["blue"]),
          ("clipboard-check", "UAT", "Business validation", C["violet"]), ("rocket", "PROD", "Approved release", C["aqua"])]
stage_html = ""
for i, (icn, t, d, col) in enumerate(stages):
    stage_html += f"""<div class="card" style="width:300px;padding:28px;text-align:center;border-top:4px solid {col}">
      <div style="display:flex;justify-content:center">{ic(icn,col,64,32)}</div>
      <div style="font-size:30px;font-weight:700;margin-top:14px">{t}</div><div style="font-size:19px;color:{C['text2']};margin-top:6px">{d}</div></div>"""
    if i < 3:
        gate = i == 2
        stage_html += f"""<div style="flex:1;display:flex;flex-direction:column;align-items:center;justify-content:center;gap:8px">
          {f'<span class="pill" style="background:{C["warning"]}22;color:{C["warning"]};border:1px solid {C["warning"]}66">{icon("lock",16,C["warning"])}Approval gate</span>' if gate else ''}
          <svg width="120" height="24"><line x1="0" y1="12" x2="104" y2="12" stroke="{C['blue_l']}" stroke-width="3"/><polygon points="104,4 120,12 104,20" fill="{C['blue_l']}"/></svg>
          <span style="font-size:17px;color:{C['muted']}">{'versioned' if not gate else 'developer requests, approver releases'}</span></div>"""
expose = [("braces", "REST API"), ("sparkles", "OpenAI-compatible API"), ("share-2", "Agent-to-Agent"), ("message-square", "Embeddable chat widget"), ("users", "Microsoft Teams")]
auto = [("calendar-clock", "Scheduled triggers"), ("folder-search", "Folder and file monitors"), ("mail", "Mailbox intake"), ("timer", "Execution timeouts"), ("bell", "Cost threshold alerts")]


def chip_row(title, items, col):
    return f"""<div class="card" style="flex:1;padding:30px 34px"><div style="font-size:27px;font-weight:600;margin-bottom:22px;color:{C['text2']}">{title}</div>
      <div style="display:flex;flex-wrap:wrap;gap:16px">{"".join(f'<span class="chip" style="font-size:24px;padding:14px 22px">{icon(i,24,col)}{t}</span>' for i,t in items)}</div></div>"""


add(
    frame(9, "Run", "From the canvas to a governed production release",
          "Every agent carries a version. Promotion to production goes through an approval, and the same agent can be exposed through several channels.",
          f"""<div style="display:flex;flex-direction:column;height:100%"><div style="display:flex;align-items:center;gap:0">{stage_html}</div>
<div style="display:flex;gap:28px;margin-top:44px;flex:1">{chip_row('Expose the agent through', expose, C['aqua'])}{chip_row('Automate and protect', auto, C['orange'])}</div></div>"""),
    "Run: Design → DEV → UAT → PROD with versioning and an approval gate before production. Expose via REST API, OpenAI-compatible API, Agent-to-Agent, embeddable chat widget, Microsoft Teams. Automate with scheduled triggers, folder monitors, mailbox intake; protect with timeouts and cost alerts.",
)

# 10 — Observe -----------------------------------------------------------------
trace_rows = [
    (0, "bot", "Agent run", C["blue"]), (1, "scroll-text", "Prompt", C["violet"]), (1, "brain", "LLM call", C["aqua"]),
    (2, "shield-check", "Guardrail check", C["aqua"]), (1, "wrench", "Tool call · database", C["orange"]),
    (1, "user-check", "Human approval", C["magenta"]), (1, "message-square", "Response", C["blue"]),
]
trace_html = "".join(
    f"""<div style="display:flex;align-items:center;gap:14px;padding:12px 16px;margin-left:{d*44}px;border-left:3px solid {c};background:rgba(255,255,255,.035);border-radius:0 12px 12px 0">
      {icon(i,22,c)}<span style="font-size:21px">{t}</span><span style="margin-left:auto;display:flex;gap:6px">{''.join(f'<span style="width:{w}px;height:8px;border-radius:4px;background:{c}88"></span>' for w in (90-d*18, 40))}</span></div>"""
    for d, i, t, c in trace_rows
)
dash = [
    ("line-chart", "Executive summary", "Runs, success rate, tokens, errors"),
    ("heart-pulse", "Infrastructure health", "Latency, 5xx, throughput, sessions"),
    ("coins", "LLM cost efficiency", "Tokens and calls by model"),
    ("gauge", "Performance", "P50, P95, P99 for runs, LLM calls and builds"),
    ("shield-alert", "Reliability", "Error classes and trends"),
    ("users", "Usage analytics", "Logins, runs by agent, tokens by model"),
]
dash_html = "".join(
    f"""<div class="card" style="padding:24px;display:flex;gap:16px;align-items:flex-start">{ic(i,C['blue'],52,26,12)}
      <div><div style="font-size:24px;font-weight:600">{t}</div><div style="font-size:18px;color:{C['text2']};margin-top:4px;line-height:1.3">{d}</div></div></div>"""
    for i, t, d in dash
)
add(
    frame(10, "Operate", "See every run, every token and every cost",
          "Traces for engineers, dashboards for operations and alerts for budget owners. Metrics flow through OpenTelemetry to Prometheus and Grafana.",
          f"""<div style="display:flex;gap:40px">
<div class="card" style="width:640px;padding:28px;flex:none"><div style="display:flex;align-items:center;gap:14px;margin-bottom:18px">{ic('scan-search',C['violet'],52,26,12)}<div style="font-size:26px;font-weight:600">Trace of a single run</div></div>
  <div style="display:flex;flex-direction:column;gap:8px">{trace_html}</div>
  <div style="font-size:18px;color:{C['muted']};margin-top:16px">Langfuse tracing, scoped per workspace. Bars illustrate relative step duration.</div></div>
<div style="flex:1;display:flex;flex-direction:column"><div style="font-size:26px;font-weight:600;margin-bottom:16px">Six Grafana dashboards out of the box</div>
  <div style="display:grid;grid-template-columns:1fr 1fr;gap:18px;flex:1">{dash_html}</div>
  <div style="display:flex;gap:12px;margin-top:22px"><span class="chip">{icon('bell',22,C['orange'])}Cost threshold alerts</span><span class="chip">{icon('flask-conical',22,C['orange'])}Evaluation datasets and evaluators</span></div></div></div>"""),
    "Operate: Langfuse traces per run scoped per workspace; OpenTelemetry metrics to Prometheus and Grafana with six dashboards: executive summary, infrastructure health, LLM cost efficiency, performance (P50/P95/P99), reliability, usage analytics. Cost threshold alerts and evaluation datasets/evaluators.",
)

# 11 — Security and enterprise readiness ---------------------------------------
sec_left = [
    ("fingerprint", "Entra ID single sign-on", "Corporate identities, linked to platform users"),
    ("key-round", "Azure Key Vault", "Provider keys and connector secrets never sit in config files"),
    ("lock", "API keys with revocation", "Token revocation and an identity blocklist"),
    ("gauge", "Rate limiting", "Per identity or IP, and per model"),
]
sec_right = [
    ("zap", "Azure Managed Redis", "Entra-authenticated cache and rate limiter"),
    ("hard-drive", "Blob storage", "Uploaded documents stay in your tenant"),
    ("globe", "Multi-region hub", "A regional gateway with circuit breakers in front of regional backends"),
    ("laptop", "Runs fully local", "The same platform runs on a single machine for demos and evaluations"),
]


def sec_col(items, col):
    return "".join(
        f"""<div class="card" style="display:flex;gap:18px;align-items:flex-start;padding:24px 26px">{ic(i,col,54,28,12)}
          <div><div style="font-size:25px;font-weight:600">{t}</div><div style="font-size:19px;color:{C['text2']};margin-top:4px;line-height:1.35">{d}</div></div></div>"""
        for i, t, d in items
    )


shield_center = f"""<div style="width:360px;flex:none;display:flex;flex-direction:column;align-items:center;justify-content:center;gap:18px;text-align:center">
  <div style="width:220px;height:220px;border-radius:50%;background:radial-gradient(circle at 50% 40%, {C['violet']}55, #121a2c 70%);border:2px solid {C['violet']};display:flex;align-items:center;justify-content:center;box-shadow:0 0 70px {C['violet']}44">{icon('shield-check',110,C['text'],1.4)}</div>
  <div style="font-size:28px;font-weight:700;line-height:1.2">Azure-native by design</div>
  <div style="font-size:19px;color:{C['text2']};line-height:1.4">Deployed inside your Azure tenant. No data leaves your subscription unless a model provider you approved is called.</div></div>"""
add(
    frame(11, "Security", "Enterprise security on your own cloud",
          "",
          f"""<div style="display:flex;gap:32px;align-items:stretch">
<div style="flex:1;display:flex;flex-direction:column;gap:18px">{sec_col(sec_left,C['blue'])}</div>{shield_center}
<div style="flex:1;display:flex;flex-direction:column;gap:18px">{sec_col(sec_right,C['aqua'])}</div></div>"""),
    "Security: Entra ID SSO, Azure Key Vault for secrets, API keys with revocation and identity blocklist, rate limiting per identity/IP and per model, Azure Managed Redis with Entra auth, Blob storage in your tenant, multi-region hub with circuit breakers, and the ability to run fully local for demos.",
)

# 12 — Use case: Email-to-Quote underwriting -----------------------------------
uw_steps = [
    ("mail-open", "Intake", "Reads the broker email and PDF attachments from the mailbox, an upload or a monitored folder."),
    ("scan-text", "Document intelligence", "Extracts 20 underwriting fields with source document, page and confidence."),
    ("clipboard-check", "Field-level validation", "Mandatory fields, formats, head-count and claims consistency. PASS, FAIL or WARN."),
    ("split", "Underwriting decision", "Proceed to quote, request clarification, or refer to an underwriter."),
    ("calculator", "Rating engine", "Deterministic premium from an editable rating table. The LLM never sets the price."),
    ("send", "Quote and dispatch", "Quote PDF and broker email, held as a draft until an underwriter approves."),
]
uw_html = ""
for i, (icn, t, d) in enumerate(uw_steps):
    uw_html += f"""<div class="card" style="flex:1;padding:24px 22px;display:flex;flex-direction:column;gap:14px;position:relative;border-top:4px solid {C['blue']}">
      <div style="display:flex;align-items:center;justify-content:space-between">{ic(icn,C['blue'],56,28)}<span style="font-size:40px;font-weight:700;color:rgba(255,255,255,.18)">{i+1}</span></div>
      <div style="font-size:27px;font-weight:600;line-height:1.2">{t}</div>
      <div style="font-size:21px;color:{C['text2']};line-height:1.4">{d}</div></div>"""
    if i < 5:
        uw_html += f'<div style="display:flex;align-items:center;flex:none"><svg width="26" height="22"><polygon points="4,2 22,11 4,20" fill="{C["blue_l"]}" fill-opacity=".8"/></svg></div>'
inputs = ["Broker email", "Employee census PDF", "Claims experience PDF", "Current policy PDF", "Company application PDF"]
outputs = ["Underwriting decision", "Validation table", "Extracted data model", "Rating breakdown", "Quote PDF", "Broker email draft", "Audit trail"]


def io_row(label, items, col, icn):
    return f"""<div style="display:flex;align-items:center;gap:18px"><span class="pill" style="background:{col}22;color:{col};border:1px solid {col}66;font-size:19px;padding:8px 16px">{icon(icn,18,col)}{label}</span>
      <div style="display:flex;gap:10px;flex-wrap:wrap">{"".join(f'<span class="chip" style="font-size:19px;padding:8px 14px">{t}</span>' for t in items)}</div></div>"""


add(
    frame(12, "Use case · Group medical insurance", "Email-to-Quote: from a broker's inbox to a governed quote",
          "An underwriting agent built entirely from AgentCore components. Documents in, an auditable decision and quote out.",
          f"""<div style="display:flex;flex-direction:column;height:100%">{io_row('Inputs', inputs, C['aqua'], 'inbox')}
<div style="display:flex;align-items:stretch;gap:0;margin:30px 0;flex:1">{uw_html}</div>
<div style="display:flex;align-items:center;gap:16px;margin-bottom:30px"><span style="font-size:22px;color:{C['text2']};flex:none;margin-right:6px">Possible decisions</span>
  <span class="pill" style="background:{C['good']}22;color:#7ee07e;border:1px solid {C['good']}66;font-size:20px;padding:9px 18px">{icon('check-circle',20,'#7ee07e')}Proceed to quote</span>
  <span class="pill" style="background:{C['warning']}22;color:{C['warning']};border:1px solid {C['warning']}66;font-size:20px;padding:9px 18px">{icon('alert-triangle',20,C['warning'])}Request clarification</span>
  <span class="pill" style="background:{C['serious']}22;color:{C['serious']};border:1px solid {C['serious']}66;font-size:20px;padding:9px 18px">{icon('user-check',20,C['serious'])}Refer to underwriter</span>
  <span style="font-size:20px;color:{C['muted']};margin-left:10px">Material validation failures put the case on hold.</span></div>
{io_row('Outputs', outputs, C['orange'], 'package-check')}</div>"""),
    "Use case: Email-to-Quote underwriting for group medical insurance. Inputs: broker email plus census, claims, policy and application PDFs. Steps: intake, document intelligence (20 fields with page and confidence), field-level validation (PASS/FAIL/WARN), underwriting decision (proceed/clarify/refer), deterministic rating engine, quote PDF and email draft. Outputs: decision, validation table, data model, rating breakdown, quote PDF, email draft, audit trail.",
)

# 13 — Use case: two submissions, safe by design ------------------------------
case_a = f"""<div class="card" style="flex:1;padding:30px;border-top:4px solid {C['good']}">
  <div style="display:flex;flex-direction:column;align-items:flex-start;gap:14px"><div style="font-size:30px;font-weight:600">Clean submission</div>
    <span class="pill" style="background:{C['good']}22;color:#7ee07e;border:1px solid {C['good']}66">{icon('check-circle',18,'#7ee07e')}Proceed to quote</span></div>
  <ul class="dots" style="--c:{C['good']};margin-top:18px">
    <li>All mandatory fields found, every one traced to a document and page</li>
    <li>Head count, dates of birth and claims totals agree across documents</li>
    <li>Within the auto-quote authority limits</li>
    <li>Premium calculated from the rating table, quote PDF generated</li>
    <li>Broker email drafted and waiting for underwriter approval</li></ul></div>"""
case_b = f"""<div class="card" style="flex:1;padding:30px;border-top:4px solid {C['warning']}">
  <div style="display:flex;flex-direction:column;align-items:flex-start;gap:14px"><div style="font-size:30px;font-weight:600">Conflicting submission</div>
    <span class="pill" style="background:{C['warning']}22;color:{C['warning']};border:1px solid {C['warning']}66">{icon('alert-triangle',18,C['warning'])}Request clarification</span></div>
  <ul class="dots" style="--c:{C['warning']};margin-top:18px">
    <li>Head count 250 in the application versus 248 in the census</li>
    <li>Dates of birth differ between census and policy for several members</li>
    <li>Claims total of 42 lakh in the summary versus 45 lakh in the detail</li>
    <li>Case placed on hold, no premium issued</li>
    <li>Clarification email drafted for the broker with every finding listed</li></ul></div>"""
principles = [
    ("calculator", "The LLM never sets the price", "Premiums come from a deterministic, editable rating table."),
    ("file-check", "Figures are protected", "Wording polish is rejected if any number changes."),
    ("user-check", "A person approves dispatch", "Nothing is sent until an authorised underwriter signs off."),
    ("scroll-text", "Every field has a source", "Document, page and confidence for each extracted value."),
]
pr_html = "".join(
    f"""<div style="display:flex;gap:16px;align-items:flex-start">{ic(i,C['violet'],52,26,12)}
      <div><div style="font-size:23px;font-weight:600">{t}</div><div style="font-size:18px;color:{C['text2']};line-height:1.35;margin-top:3px">{d}</div></div></div>"""
    for i, t, d in principles
)
add(
    frame(13, "Use case · What the client sees", "Same agent, two submissions, two defensible outcomes",
          "",
          f"""<div style="display:flex;gap:28px;align-items:stretch;height:100%">
<div style="flex:1.6;display:flex;gap:28px">{case_a}{case_b}</div>
<div class="card" style="flex:.9;padding:30px;display:flex;flex-direction:column;gap:30px;border-top:4px solid {C['violet']}"><div style="font-size:28px;font-weight:600">Safe by design</div>{pr_html}</div></div>""",
          extra_css="ul.dots li{font-size:24px;margin-top:16px}"),
    "Two sample submissions. Clean: all fields found and consistent, within authority, premium from rating table, quote PDF and email draft pending approval. Conflicting: head count 250 vs 248, DOB mismatches, claims 42 lakh vs 45 lakh, case on hold, clarification email drafted. Principles: LLM never sets the price, figures protected, human approves dispatch, every field has a source.",
)

# 14 — Engagement path --------------------------------------------------------
path = [
    ("laptop", C["blue"], "Demonstrate",
     ["Platform running on a single machine, no cloud dependency", "Underwriting agent on the bundled samples and your own documents",
      "Hands-on session with your business and IT teams", "Review of security, governance and integration questions"],
     "A clear go or no-go on your first use case"),
    ("building-2", C["violet"], "Pilot",
     ["Deployed in your Azure tenant with Entra ID single sign-on", "One business process built end to end with your experts",
      "Governance, guardrails and tracing switched on from day one", "Baseline measured: turnaround time, touch time, error rate"],
     "One process live with measured results"),
    ("trending-up", C["aqua"], "Scale",
     ["Catalogue of approved agents, models, tools and guardrails", "New processes reuse the same components and roles",
      "Multi-region deployment and cost controls", "Evaluation datasets drive continuous improvement"],
     "A governed agent estate across the business"),
]
path_html = ""
for i, (icn, col, t, bs, outcome) in enumerate(path):
    path_html += f"""<div class="card" style="flex:1;padding:36px 34px;border-top:4px solid {col};display:flex;flex-direction:column">
      <div style="display:flex;align-items:center;gap:18px">{ic(icn,col,68,34)}<div><div style="font-size:19px;letter-spacing:.16em;text-transform:uppercase;color:{C['muted']}">Stage {i+1}</div><div style="font-size:36px;font-weight:700">{t}</div></div></div>
      <ul class="dots" style="--c:{col};margin-top:22px">{"".join(f"<li>{b}</li>" for b in bs)}</ul>
      <div style="margin-top:auto;padding-top:22px;border-top:1px solid rgba(255,255,255,.1);display:flex;gap:12px;align-items:center;font-size:22px;color:{C['text2']};line-height:1.3">{icon('flag',24,col)}<span><b style="color:{C['text']}">Outcome:</b> {outcome}</span></div></div>"""
    if i < 2:
        path_html += f'<div style="display:flex;align-items:center;padding:0 10px"><svg width="44" height="40"><polygon points="6,4 38,20 6,36" fill="{C["blue_l"]}" fill-opacity=".8"/></svg></div>'
add(
    frame(14, "Engagement path", "From demonstration to a governed agent estate",
          "Each stage reuses everything built in the previous one. Agents, roles, models and guardrails carry over unchanged.",
          f'<div style="display:flex;align-items:stretch;height:100%">{path_html}</div>',
          extra_css="ul.dots li{font-size:24px;margin-top:14px}"),
    "Engagement path. Demonstrate: platform on a single machine, underwriting agent on sample and client documents. Pilot: deployed in the client's Azure tenant with SSO, one process end to end, governance on. Scale: catalogue of approved agents, multi-region, cost controls, evaluation.",
)

# 15 — Why AgentCore -----------------------------------------------------------
why = [
    ("sparkles", C["blue"], "Model-agnostic", "Switch providers behind the gateway without touching a single agent.", "7 provider families, one endpoint"),
    ("shield-check", C["violet"], "Governance by design", "Roles, approvals, guardrails and audit are part of the runtime.", "Approvals for models, tools and releases"),
    ("mail", C["aqua"], "Microsoft 365-native", "Outlook, SharePoint and Teams are first-class components.", "Microsoft Graph and Entra ID throughout"),
    ("rocket", C["orange"], "Production-grade", "Versioned releases, approval gates, rate limits and timeouts.", "DEV, UAT and PROD environments"),
    ("activity", C["magenta"], "Fully observable", "Traces, metrics and cost per agent and per model.", "6 Grafana dashboards, Langfuse traces"),
    ("file-check", C["yellow"], "Proven on a real process", "A complete underwriting agent ships with the platform.", "6-step pipeline, 2 worked cases"),
]
why_html = "".join(
    f"""<div class="card" style="padding:32px;display:flex;flex-direction:column;gap:16px;border-top:4px solid {c}">
      {ic(i,c,62,32)}<div style="font-size:31px;font-weight:600">{t}</div><div style="font-size:23px;color:{C['text2']};line-height:1.45">{d}</div>
      <div style="margin-top:auto;display:flex;align-items:center;gap:10px;font-size:20px;color:{c};font-weight:600">{icon('check',20,c,2.5)}{p}</div></div>"""
    for i, c, t, d, p in why
)
add(
    frame(15, "Summary", "Why AgentCore", "",
          f'<div style="display:grid;grid-template-columns:repeat(3,1fr);gap:26px;height:100%">{why_html}</div>'),
    "Why AgentCore: model-agnostic, governance by design, Microsoft 365-native, production-grade runtime, fully observable, proven on a real underwriting process.",
)

# 16 — Closing -----------------------------------------------------------------
next_steps = [
    ("search", "Choose a process", "Documents, decisions and a mailbox: that is where agents pay off first."),
    ("folder-open", "Share sample data", "A handful of real cases and the rules your experts apply today."),
    ("play", "Hands-on demonstration", "Your use case running on AgentCore, with your team at the keyboard."),
]
ns_html = "".join(
    f"""<div style="display:flex;gap:20px;align-items:flex-start"><div style="width:52px;height:52px;border-radius:50%;background:{C['blue']};display:flex;align-items:center;justify-content:center;font-size:24px;font-weight:700;flex:none">{i+1}</div>
      <div><div style="font-size:28px;font-weight:600">{t}</div><div style="font-size:21px;color:{C['text2']};line-height:1.4;margin-top:4px">{d}</div></div></div>"""
    for i, (_, t, d) in enumerate(next_steps)
)
add(
    bare(f"""
<div style="position:absolute;right:-60px;bottom:-80px">{network_svg(1100, 760, .45)}</div>
<div style="position:absolute;left:120px;top:140px;display:flex;align-items:center;gap:22px">{logo(72)}<span style="font-size:52px;font-weight:700;letter-spacing:-.02em">AgentCore</span></div>
<div style="position:absolute;left:120px;top:290px;font-size:84px;font-weight:700;line-height:1.05;letter-spacing:-.02em;max-width:1000px">Let's build your<br/>first agent together</div>
<div style="position:absolute;left:120px;top:560px;display:flex;flex-direction:column;gap:34px;max-width:900px">{ns_html}</div>
<div style="position:absolute;right:120px;bottom:80px;color:{C['muted']};font-size:22px;text-align:right">Thank you</div>
"""),
    "Closing. Next steps: choose a process with documents, decisions and a mailbox; share sample data and the rules experts apply; a hands-on demonstration of the client's use case on AgentCore.",
)


# ----------------------------------------------------------------------------
# Render and assemble
# ----------------------------------------------------------------------------
def render(i: int, html: str) -> pathlib.Path:
    html_path = HTML_DIR / f"slide_{i:02d}.html"
    png_path = PNG_DIR / f"slide_{i:02d}.png"
    html_path.write_text(html, encoding="utf-8")
    cmd = [
        str(CHROME), "--headless=new", "--disable-gpu", "--hide-scrollbars", "--no-first-run",
        "--disable-extensions", "--force-device-scale-factor=2", "--window-size=1920,1080",
        "--virtual-time-budget=2500", f"--screenshot={png_path}", html_path.as_uri(),
    ]
    subprocess.run(cmd, check=True, capture_output=True)
    return png_path


def main(only: list[int] | None = None) -> None:
    HTML_DIR.mkdir(parents=True, exist_ok=True)
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    pngs = []
    for i, (html, _) in enumerate(SLIDES, start=1):
        if only and i not in only:
            pngs.append(PNG_DIR / f"slide_{i:02d}.png")
            continue
        pngs.append(render(i, html))
        print(f"rendered slide {i:02d}")
    if _missing_icons:
        print("MISSING ICONS (fell back to circle):", ", ".join(sorted(_missing_icons)))
    if only:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for png, (_, notes) in zip(pngs, SLIDES):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        s.notes_slide.notes_text_frame.text = notes
    prs.save(PPTX_PATH)
    print(f"wrote {PPTX_PATH} with {len(SLIDES)} slides")


if __name__ == "__main__":
    only_arg = [int(a) for a in sys.argv[1:]] if len(sys.argv) > 1 else None
    main(only_arg)
